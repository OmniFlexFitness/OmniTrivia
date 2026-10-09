"""
Builds the question template workbook, `questions.template.xlsx`.

    pip install openpyxl          # once
    python3 scripts/generate-template.py
    python3 scripts/generate-template.py --preview   # and redraw its picture

The first tab, Questions, is `questions.example.csv` cell for cell: one row of
every question format, ready to overwrite. The tabs after it explain how to
fill it in. The output is committed, so nobody needs to run this to use the
template, only to rebuild it after the example file or the formats change.
`npm run check-question-types` fails when the committed workbook no longer
matches `questions.example.csv`.

Why the Questions tab has to come first: the importer reads one tab. Upload the
workbook to Google Drive and paste its share link, and with no `gid=` in that
link the app reads the leftmost tab. Saving to CSV from Excel saves the open
tab, and Questions is the tab the workbook opens on. Nothing but questions can
go on that tab either, since every row on it is read as one: the notes live in
cell comments and on the other tabs.

`--preview` also draws `docs/images/question-format/template.png`, the
Questions tab as a spreadsheet program shows it, for QUESTION_FORMAT.md. It
needs LibreOffice (`soffice`) and Poppler (`pdftoppm`) on the PATH, and Pillow.
"""

import csv
import io
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "questions.example.csv"
OUTPUT = ROOT / "questions.template.xlsx"
PREVIEW = ROOT / "docs" / "images" / "question-format" / "template.png"
GUIDE_URL = "https://github.com/OmniFlexFitness/OmniTrivia/blob/master/QUESTION_FORMAT.md"

# A fixed date for everything that would otherwise record when the file was
# built, so building twice from the same sources gives the same bytes and a
# rebuild that changes nothing shows no diff.
STAMP = datetime(2026, 1, 1)

FONT = "Arial"
INK = "1E293B"  # slate-800: header fill and body text
NEON = "00D4FF"  # the app's neon blue, as a rule under the header
PAPER = "F8FAFC"  # slate-50: the reference tabs' banding

# Rows the template leaves room for. Validation and the type colours cover
# this many, which is several nights of questions.
ROWS = 1000

# Every format, with the badge colour the review screen gives it (the
# Tailwind 100 shade of the same hue, so dark text stays readable).
TYPES = [
    ("MULTIPLE_CHOICE", "E0F2FE"),  # sky
    ("MULTI_SELECT", "CFFAFE"),  # cyan
    ("TRUE_FALSE", "F3E8FF"),  # purple
    ("TYPE_ANSWER", "FEF9C3"),  # yellow
    ("SLIDER", "FFEDD5"),  # orange
    ("RANGE", "FEF3C7"),  # amber
    ("NUMBER", "ECFCCB"),  # lime
    ("PIN", "D1FAE5"),  # emerald
    ("PUZZLE", "FEE2E2"),  # red
    ("MATCH", "FCE7F3"),  # pink
    ("CATEGORIZE", "FAE8FF"),  # fuchsia
    ("SCRAMBLE", "EDE9FE"),  # violet
]
MARGINS = ["none", "low", "medium", "high", "maximum"]

# What each column is for, as a note on its header cell. Hover the header in
# Excel, or in Google Sheets once the file is opened there, to read it.
HEADER_NOTES = {
    "type": "Which format the row is. Pick one from the list. Leave it blank for "
    "multiple choice, or for the shorthand: a row with no options becomes a "
    "typed answer, or true/false when the answer is True or False. Spellings "
    "like True/False, TF, Select all that apply or Puzzle · order work too, but "
    "a word the game does not know leaves the row out as an unknown format.",
    "category": "Groups questions into rounds: one category per round on the wheel. "
    "Science, History, Geography, Pop Culture, Sports, Tech, Art, Literature, "
    "Music and Food get their own icon and colour; any other name works too.",
    "question": "What players read. Keep it on one line: no line breaks in any cell.",
    "option1": "Options 1 to 6 mean different things for different formats. See the "
    "Question types tab. Add more option columns (option7, option8...) if a "
    "question needs them; leave unused ones empty.",
    "correctAnswer": "The answer. Its form depends on the format: one option, several "
    "joined with |, True or False, a number, or a word to unscramble. See the "
    "Question types tab.",
    "explanation": "Optional. Shown with the answer after the round.",
    "image": "Optional on any format, required for PIN: an https:// link to a "
    "picture, or world / usa for a built-in map.",
    "margin": "SLIDER, NUMBER and PIN only: how much a near miss still scores. "
    "none, low, medium, high or maximum.",
    "unit": "SLIDER, RANGE and NUMBER only: what the number counts, shown after it "
    "(km, m, °F, or years for a span of years). Leave it blank for a calendar "
    "year, or 1969 shows as 1,969 years.",
    "timeLimit": "Optional: seconds on the clock for this question, 5 to 120. Blank "
    "uses the format's own clock.",
}

# One row per format for the Question types tab.
TYPE_GUIDE_HEADER = [
    "type",
    "What players do",
    "option1 … option6",
    "correctAnswer",
    "Other columns",
    "Clock",
    "Scoring",
    "Other spellings that work",
]
TYPE_GUIDE = [
    ("MULTIPLE_CHOICE", "Tap one option", "The choices: two or more",
     "The right option, written as it is in its column (not its letter or number)",
     "—", "15s", "All or nothing", "MC, Quiz, Choice, Single select, or leave blank"),
    ("MULTI_SELECT", "Tick every right option", "The choices: three or more",
     "Every right option, joined with | or ; (Helium|Argon)", "—", "20s",
     "A wrong tick cancels a right one", "Select all that apply, Multiple select, Checkboxes"),
    ("TRUE_FALSE", "Tap True or False", "True and False, or leave both empty",
     "True or False (T/F, Yes/No, Fact/Fiction and 1/0 work too)", "—", "12s",
     "All or nothing", "TF, True/False, True or False, Fact or fiction"),
    ("TYPE_ANSWER", "Type the answer",
     "Every spelling you will accept. The first is the one the reveal shows",
     "Required, but not graded once the options are filled in: repeat option1. "
     "With none, the answer and each part of a comma list are accepted",
     "—", "20s",
     "All or nothing. One accepted spelling, or several in any order (a, b and c) when every one is accepted; "
     "case, curly quotes and extra spaces ignored, spelling and accents not",
     "Typed answer, Short answer, Fill in the blank, Free text"),
    ("SLIDER", "Slide to a number",
     "Lowest, highest, step, then the lowest and highest answer that score in full",
     "Required, but not graded or shown when option4 and option5 are filled in: the reveal shows that band. "
     "With only three options: the answer, or a band like 1968-1970 with nothing after it (a unit goes in the unit column)",
     "margin, unit", "20s", "Inside the target scores in full; margin pays a near miss",
     "Slide, Scale"),
    ("RANGE", "Drag two handles to catch the answer", "Lowest, highest, step",
     "The exact answer, or a band like 8800-8900 with nothing after it (a unit goes in the unit column)",
     "unit", "20s",
     "Caught inside a tenth of the scale scores in full; wider pays less, past half pays nothing",
     "Range slider, Bracket, Between"),
    ("NUMBER", "Type a number, no scale shown", "Optional: a tolerance that still counts as dead on",
     "The number (commas and a unit are fine: 384,400 km)", "margin, unit", "20s",
     "Exact or within the margin scores; margin is medium unless set", "Closest, Closest guess, Estimate, Guess"),
    ("PIN", "Drop a pin on a map or a picture",
     "Map: latitude, longitude, km that still score in full. Picture: % across, % down, radius %",
     "What the answer key calls the spot", "image (world, usa or https://…), margin",
     "25s", "Inside the ring scores in full; margin pays a near miss",
     "Pin answer, Drop pin, Map, Hotspot"),
    ("PUZZLE", "Put the items in order", "The items, in the right order: two or more, all different",
     "Required, but not graded when the options are filled in: write the order again. "
     "With none: the order, joined with | (the type column cannot be blank)", "—", "30s",
     "All or nothing", "Puzzle · order, Order, Sequence, Rank, Timeline"),
    ("MATCH", "Pair each item with its partner",
     "One pair per column: item = partner (-> → => :: work too). Two or more pairs",
     "Ignored when the options are filled in: write See the pairs", "—", "30s",
     "Pays for every pair right beyond guessing", "Puzzle · match, Matching, Pairs, Connect"),
    ("CATEGORIZE", "Sort items into groups",
     "One item per column: item = group. Two to four groups, spelled the same every time",
     "Ignored when the options are filled in: write See the groups", "—", "30s",
     "Pays for every item right beyond guessing", "Puzzle · sort, Sort, Group, Buckets, Classify"),
    ("SCRAMBLE", "Rebuild a word from shuffled letters", "Leave empty: they are ignored",
     "The word or phrase, 3 to 20 letters and digits, no brackets. The question is the clue",
     "—", "25s", "All or nothing", "Puzzle · unscramble, Unscramble, Anagram, Jumble"),
    ("(blank)", "Shorthand: no type, no options", "Leave every option empty",
     "True or False makes a true/false question; anything else a typed answer", "—",
     "12s / 20s", "As true/false or typed answer", "—"),
]

HOW_TO = [
    ("Start here", None),
    ("1. Write your questions on the Questions tab, one per row. Overwrite the example rows; keep row 1, the headings, exactly as it is.", None),
    ("2. Every white cell under the headings is yours to edit. The coloured cell in the type column just shows which format the row is.", None),
    ("3. The type column has a drop-down of the twelve formats. The Question types tab says what goes in the option and correctAnswer columns for each one.", None),
    ("4. Hover a heading for a note on what that column is for.", None),
    ("", None),
    ("Getting it into OmniTrivia", None),
    ("From Excel: File → Save As → CSV UTF-8 (Comma delimited), with the Questions tab open. Then HOST GAME → IMPORT MY OWN QUESTIONS → Select File. Plain \"CSV (Comma delimited)\" mangles dashes, accents and °.", None),
    ("From Google Drive: upload this file, set sharing to \"Anyone with the link can view\", and paste the share link on the import screen. The app reads the first tab, which is Questions.", None),
    ("In Google Sheets: File → Save as Google Sheets, then share and paste the link the same way. If you import it into a spreadsheet you already have (File → Import → Insert new sheet(s)), open the Questions tab and paste the URL from the address bar (it ends in gid=…): a plain share link reads that spreadsheet's original tab.", None),
    ("", None),
    ("Rules that catch people out", None),
    ("• No line breaks inside a cell. A line break splits the row in two.", None),
    ("• Keep the Questions tab first, and put nothing on it but questions: every row on it is read as one.", None),
    ("• A row that cannot be played is left out. The review screen says how many and why (the first few row numbers are in the browser's console), so check it shows the number of questions you expect.", None),
    ("• Pictures load straight from their link on every phone, so use a host that allows it, and check them on the review screen.", None),
    ("", None),
    ("The full guide, with a picture of every format:", None),
    (GUIDE_URL, GUIDE_URL),
]


def cell_value(text: str):
    """A number where the spreadsheet would show one, so Excel does not flag
    "number stored as text" on every option of a slider. Anything else, a
    range like 1968-1970 included, stays exactly as written."""
    # Only when the number reads back exactly as written: 0.10 would come back
    # out of a spreadsheet as 0.1, and 007 as 7.
    if re.fullmatch(r"-?[1-9]\d*|0", text):
        return int(text)
    if re.fullmatch(r"-?\d+\.\d+", text) and repr(float(text)) == text:
        return float(text)
    return text


def header_style(cell) -> None:
    cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", fgColor=INK)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = Border(bottom=Side(style="medium", color=NEON))


def build_questions(ws) -> None:
    with SOURCE.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    header, body = rows[0], rows[1:]

    ws.append(header)
    for row in body:
        ws.append([cell_value(value) if value != "" else None for value in row])

    body_font = Font(name=FONT, size=10, color=INK)
    for row in ws.iter_rows(min_row=2, max_row=len(body) + 1):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    widths = {
        "type": 18, "category": 13, "question": 44, "correctAnswer": 26,
        "explanation": 48, "image": 10, "margin": 9, "unit": 7, "timeLimit": 9,
    }
    for index, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=index)
        header_style(cell)
        note = HEADER_NOTES.get(name) or (
            HEADER_NOTES["option1"] if name.startswith("option") else None
        )
        if note:
            comment = Comment(note, "OmniTrivia")
            comment.width, comment.height = 320, 150
            cell.comment = comment
        ws.column_dimensions[get_column_letter(index)].width = widths.get(
            name, 16 if name.startswith("option") else 14
        )
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "D2"

    # Printed, the whole sheet fits across one landscape page, headings
    # repeated on every page.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    column = {name: get_column_letter(i) for i, name in enumerate(header, start=1)}

    # The format drop-down. A warning rather than a refusal, because the
    # importer reads plenty of other spellings and a host who types TF should
    # not be told no.
    type_list = DataValidation(
        type="list",
        formula1='"' + ",".join(name for name, _ in TYPES) + '"',
        allow_blank=True,
        showErrorMessage=True,
        showInputMessage=True,
        errorStyle="warning",
        errorTitle="Not one of the twelve names",
        error="Other spellings like TF or Select all that apply still import. "
        "Check the Question types tab.",
        promptTitle="Format",
        prompt="Pick a format, or leave blank for multiple choice.",
    )
    type_list.add(f"{column['type']}2:{column['type']}{ROWS}")
    ws.add_data_validation(type_list)

    margin_list = DataValidation(
        type="list",
        formula1='"' + ",".join(MARGINS) + '"',
        allow_blank=True,
        showErrorMessage=True,
        showInputMessage=True,
        errorStyle="warning",
        error="margin is none, low, medium, high or maximum.",
        promptTitle="Margin",
        prompt="SLIDER, NUMBER and PIN: how much a near miss scores.",
    )
    margin_list.add(f"{column['margin']}2:{column['margin']}{ROWS}")
    ws.add_data_validation(margin_list)

    clock = DataValidation(
        type="whole",
        operator="between",
        formula1="5",
        formula2="120",
        allow_blank=True,
        showErrorMessage=True,
        showInputMessage=True,
        errorStyle="warning",
        error="The clock runs from 5 to 120 seconds.",
        promptTitle="Seconds",
        prompt="Blank uses the format's own clock.",
    )
    clock.add(f"{column['timeLimit']}2:{column['timeLimit']}{ROWS}")
    ws.add_data_validation(clock)

    # Each format's type cell in its review-screen colour, so a sheet of mixed
    # formats can be read down the first column.
    first = column["type"]
    for name, colour in TYPES:
        ws.conditional_formatting.add(
            f"{first}2:{first}{ROWS}",
            FormulaRule(
                formula=[f'UPPER(TRIM({first}2))="{name}"'],
                fill=PatternFill("solid", fgColor=colour, bgColor=colour),
                font=Font(name=FONT, bold=True, color=INK),
            ),
        )


def build_reference(ws, header, rows, widths) -> None:
    ws.append(header)
    for index in range(1, len(header) + 1):
        header_style(ws.cell(row=1, column=index))
    for number, row in enumerate(rows, start=2):
        ws.append(list(row))
        for cell in ws[number]:
            cell.font = Font(name=FONT, size=10, color=INK, bold=cell.column == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if number % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=PAPER)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "B2"


def build_how_to(ws) -> None:
    ws.column_dimensions["A"].width = 110
    headings = {"Start here", "Getting it into OmniTrivia", "Rules that catch people out"}
    for number, (text, link) in enumerate(HOW_TO, start=1):
        cell = ws.cell(row=number, column=1, value=text)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if text in headings:
            cell.font = Font(name=FONT, size=12, bold=True, color=INK)
        elif link:
            cell.hyperlink = link
            cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        else:
            cell.font = Font(name=FONT, size=10, color=INK)


def main() -> None:
    workbook = Workbook()
    questions = workbook.active
    questions.title = "Questions"
    build_questions(questions)

    build_how_to(workbook.create_sheet("How to use"))
    build_reference(
        workbook.create_sheet("Question types"),
        TYPE_GUIDE_HEADER,
        TYPE_GUIDE,
        [17, 24, 34, 38, 24, 8, 30, 30],
    )

    workbook.active = 0
    workbook.properties.created = STAMP
    workbook.properties.creator = "OmniTrivia"
    workbook.properties.lastModifiedBy = "OmniTrivia"
    workbook.properties.title = "OmniTrivia question template"
    save_reproducibly(workbook, OUTPUT)
    print(f"Wrote {OUTPUT.relative_to(ROOT)}")
    if "--preview" in sys.argv:
        draw_preview()


def draw_preview() -> None:
    """The Questions tab, printed by LibreOffice and cropped: the picture at
    the top of QUESTION_FORMAT.md's template section. Only the columns from
    type to correctAnswer, so the text is big enough to read on a page."""
    from openpyxl import load_workbook
    from PIL import Image, ImageChops

    for tool in ("soffice", "pdftoppm"):
        if not shutil.which(tool):
            sys.exit(f"--preview needs {tool} on the PATH")
    with tempfile.TemporaryDirectory() as scratch:
        work = Path(scratch)
        workbook = load_workbook(OUTPUT)
        sheet = workbook["Questions"]
        sheet.print_area = f"A1:J{sheet.max_row}"
        sheet.print_title_rows = None
        for other in workbook.sheetnames[1:]:
            del workbook[other]
        workbook.save(work / "preview.xlsx")
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(work), str(work / "preview.xlsx")],
            check=True,
            capture_output=True,
        )
        subprocess.run(
            ["pdftoppm", "-png", "-r", "170", "-f", "1", "-l", "1", str(work / "preview.pdf"), str(work / "page")],
            check=True,
        )
        page = Image.open(next(work.glob("page*.png"))).convert("RGB")
        # Trim the paper margin, keeping a little of it.
        blank = Image.new("RGB", page.size, (255, 255, 255))
        box = ImageChops.difference(page, blank).getbbox()
        left, top, right, bottom = box
        page = page.crop((max(0, left - 16), max(0, top - 16), right + 16, bottom + 16))
        PREVIEW.parent.mkdir(parents=True, exist_ok=True)
        page.save(PREVIEW, optimize=True)
    print(f"Wrote {PREVIEW.relative_to(ROOT)}")


def save_reproducibly(workbook, path: Path) -> None:
    """Save, then re-pack the archive with fixed dates.

    openpyxl stamps the save time into docProps/core.xml and every zip entry,
    and offers no way not to, so both are put back to `STAMP` afterwards."""
    buffer = io.BytesIO()
    workbook.save(buffer)
    stamp = STAMP.strftime("%Y-%m-%dT%H:%M:%SZ")
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(
        path, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*",
                    rb"\g<1>" + stamp.encode(),
                    data,
                )
            fixed = zipfile.ZipInfo(entry.filename, date_time=STAMP.timetuple()[:6])
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = 0o644 << 16
            target.writestr(fixed, data)


if __name__ == "__main__":
    main()
